import argparse
import os
import random
import re
import shutil
import string
import subprocess
import sys
import tempfile
import threading
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from typing import List, Optional
from urllib.parse import unquote, urljoin, urlparse

import requests
from openpyxl import Workbook

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext, ttk
except ImportError:
    tk = None

BASE_URL = "https://e-consultaruc.sunat.gob.pe/cl-ti-itmrconsmulruc/jrmS00Alias"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "es-PE,es;q=0.9,en;q=0.8",
    "Referer": BASE_URL,
}


def parse_cell_to_string(cell) -> str:
    """Limpia celdas con interlineado, notación científica y decimales de Excel."""
    if cell is None:
        return ""
    if isinstance(cell, (int, float)):
        if isinstance(cell, float):
            if cell.is_integer():
                return str(int(cell))
            return str(int(float(cell)))
        return str(cell)
    
    val = str(cell).strip()
    val = re.sub(r'[\r\n]+', ' ', val).strip()
    
    if re.match(r'^[+-]?\d+(\.\d+)?[eE][+-]?\d+$', val) or re.match(r'^[+-]?\d+\.0+$', val):
        try:
            f_val = float(val)
            return str(int(f_val))
        except (ValueError, OverflowError):
            pass
            
    return val


def generate_sunat_token(length: int = 52) -> str:
    alphabet = string.ascii_lowercase + string.digits
    token = ""
    while len(token) < length:
        token += ''.join(random.choice(alphabet) for _ in range(16))
    return token[:length]


def normalize_ruc_text(value: str) -> str:
    value = value.strip()
    if not value:
        return ""
    cleaned = re.sub(r"[^0-9]", "", value)
    if len(cleaned) < 11 and len(cleaned) >= 8:
        cleaned = cleaned.zfill(11)
    return cleaned if cleaned.endswith("|") else cleaned + "|"


def extract_document_from_row(row: tuple, row_index: int) -> Optional[str]:
    if not row:
        return None

    cells = []
    for cell in row:
        c_str = parse_cell_to_string(cell)
        if c_str:
            cells.append(c_str)
            
    if not cells:
        return None

    if row_index == 1:
        header_text = " ".join(cells).lower()
        if any(keyword in header_text for keyword in ["num", "documento", "tipo", "nombre", "apellido", "ruc", "dni"]):
            return None

    for cell in cells:
        parsed_cell = parse_cell_to_string(cell)
        cleaned = re.sub(r"[^0-9]", "", parsed_cell)
        if len(cleaned) == 11:
            return normalize_ruc_text(cleaned)

    return None


class SunatHtmlTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: List[List[List[str]]] = []
        self.current_table: List[List[str]] = []
        self.current_row: List[str] = []
        self.current_cell: List[str] = []
        self.in_table = False
        self.in_cell = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        if tag == "table":
            self.in_table = True
            self.current_table = []
        elif self.in_table and tag == "tr":
            self.current_row = []
        elif self.in_table and tag in ("td", "th"):
            self.in_cell = True
            self.current_cell = []

    def handle_endtag(self, tag: str) -> None:
        if self.in_table and tag in ("td", "th"):
            self.in_cell = False
            cell_text = "".join(self.current_cell).strip()
            self.current_row.append(cell_text)
        elif self.in_table and tag == "tr":
            if self.current_row:
                self.current_table.append(self.current_row)
        elif tag == "table" and self.in_table:
            self.in_table = False
            if self.current_table:
                self.tables.append(self.current_table)

    def handle_data(self, data: str) -> None:
        if self.in_cell:
            self.current_cell.append(data)


def parse_html_table(html: str) -> List[List[str]]:
    parser = SunatHtmlTableParser()
    parser.feed(html)
    parser.close()
    return parser.tables[0] if parser.tables else []


def parse_pipe_delimited_text(text: str) -> List[List[str]]:
    rows: List[List[str]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip("\r")
        if not line:
            continue
        rows.append(line.split("|"))
    return rows


def get_fallback_output_dir() -> str:
    fallback_dir = os.path.join(os.getcwd(), "resultados_sunat")
    os.makedirs(fallback_dir, exist_ok=True)
    return fallback_dir


def get_filename_from_response(response: requests.Response) -> Optional[str]:
    header = response.headers.get("content-disposition", "")
    if not header:
        return None
    filename_match = re.search(r"filename\*=UTF-8''(?P<name>[^;]+)", header)
    if filename_match:
        return unquote(filename_match.group("name"))
    filename_match = re.search(r"filename=(?P<name>\"?[^\";]+\"?)", header)
    if filename_match:
        return filename_match.group("name").strip('"')
    return None


def is_archive_content(response: requests.Response) -> bool:
    content_type = response.headers.get("content-type", "").lower()
    if "rar" in content_type or "zip" in content_type:
        return True
    content_start = response.content[:8]
    return content_start.startswith(b"Rar!") or content_start.startswith(b"PK\x03\x04")


def save_binary_response(response: requests.Response, save_path: str) -> str:
    save_path = os.path.normpath(save_path)
    with open(save_path, "wb") as f:
        f.write(response.content)
    return save_path


def find_archive_link(html: str) -> Optional[str]:
    matches = re.findall(r'href=["\']([^"\']+\.(?:zip|rar)(?:\?[^"\']*)?)["\']', html, flags=re.IGNORECASE)
    if matches:
        return matches[0]
    
    matches_text = re.findall(r'<a\s+[^>]*href=["\']([^"\']+)["\'][^>]*>(?:[^<]*\.(?:zip|rar))</a>', html, flags=re.IGNORECASE)
    if matches_text:
        return matches_text[0]
        
    matches_download = re.findall(r'href=["\']([^"\']*(?:descargaArchivo|download|descarga)[^"\']*)["\']', html, flags=re.IGNORECASE)
    if matches_download:
        return matches_download[0]

    return None


def get_remote_filename_from_url(url: str) -> str:
    parsed = urlparse(url)
    name = os.path.basename(parsed.path)
    if not name:
        name = "resultado_sunat"
    return name


def save_response_to_excel(text: str, save_path: str) -> str:
    normalized = text.strip()
    if normalized.lower().startswith("<html") or normalized.lower().startswith("<!doctype html"):
        rows = parse_html_table(normalized)
    elif "|" in normalized:
        rows = parse_pipe_delimited_text(normalized)
    else:
        rows = [[normalized]]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SUNAT"

    if not rows:
        sheet.append([normalized])
    else:
        for row in rows:
            sheet.append(row)

    save_path = os.path.normpath(save_path)
    try:
        workbook.save(save_path)
        return save_path
    except PermissionError:
        fallback_dir = get_fallback_output_dir()
        fallback_path = os.path.join(fallback_dir, os.path.basename(save_path))
        workbook.save(fallback_path)
        return fallback_path
    except OSError as exc:
        if getattr(exc, "errno", None) == 13:
            fallback_dir = get_fallback_output_dir()
            fallback_path = os.path.join(fallback_dir, os.path.basename(save_path))
            workbook.save(fallback_path)
            return fallback_path
        raise


def download_archive_from_html(response: requests.Response, output_dir: str) -> Optional[str]:
    if not response.text:
        return None
    archive_link = find_archive_link(response.text)
    if not archive_link:
        return None

    archive_url = urljoin(response.url, archive_link)
    archive_response = requests.get(archive_url, headers=HEADERS, timeout=120)
    archive_response.encoding = archive_response.apparent_encoding
    archive_response.raise_for_status()

    filename = get_filename_from_response(archive_response) or get_remote_filename_from_url(archive_url)
    _, ext = os.path.splitext(filename)
    if not ext:
        ext = ".zip" if is_archive_content(archive_response) else ".txt"
        filename += ext
    save_path = os.path.join(output_dir, filename)
    return save_binary_response(archive_response, save_path)


def save_sunat_response(response: requests.Response, output_dir: str, batch_index: int) -> str:
    if response.text and ("cuerpo" in response.text or "La aplicación ha retornado" in response.text):
        if "no son válidos o no existen" in response.text or "no son validos o no existen" in response.text:
            filename = f"resultado_sunat_batch_{batch_index}_ERROR_rucs_invalidos.txt"
            save_path = os.path.normpath(os.path.join(output_dir, filename))
            with open(save_path, "w", encoding="utf-8") as f:
                f.write("ERROR DE SUNAT: Los números de RUC ingresados no son válidos o no existen.\n")
            return save_path

        error_match = re.search(r'<p class="error">(.*?)</p>', response.text, flags=re.DOTALL)
        if error_match:
            err_msg = error_match.group(1).strip()
            err_msg = err_msg.replace("&amp;", "&").replace("&oacute;", "ó").replace("&iacute;", "í").replace("&aacute;", "á").replace("&eacute;", "é").replace("&uacute;", "ú").replace("&ntilde;", "ñ")
            filename = f"resultado_sunat_batch_{batch_index}_ERROR.txt"
            save_path = os.path.normpath(os.path.join(output_dir, filename))
            with open(save_path, "w", encoding="utf-8") as f:
                f.write(f"ERROR DE SUNAT: {err_msg}\n")
            return save_path

    filename = get_filename_from_response(response) or f"resultado_sunat_batch_{batch_index}"
    _, ext = os.path.splitext(filename)
    ext = ext.lower()

    if not ext:
        ext = ".rar" if is_archive_content(response) else ".txt"
        filename += ext

    save_path = os.path.join(output_dir, filename)

    if is_archive_content(response):
        return save_binary_response(response, save_path)

    archive_from_html = download_archive_from_html(response, output_dir)
    if archive_from_html:
        return archive_from_html

    if ext == ".txt" or response.headers.get("content-type", "").lower().startswith("text"):
        save_path = os.path.normpath(save_path)
        with open(save_path, "w", encoding="utf-8", errors="replace") as f:
            f.write(response.text)
        return save_path

    return save_response_to_excel(response.text, os.path.join(output_dir, f"resultado_sunat_batch_{batch_index}.xlsx"))


def load_rucs_from_file(file_path: str) -> List[str]:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"No existe el archivo: {file_path}")

    if file_path.lower().endswith(".txt"):
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            values = [normalize_ruc_text(parse_cell_to_string(line)) for line in f if line.strip()]
    elif file_path.lower().endswith((".xls", ".xlsx")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError(
                "Para abrir archivos Excel es necesario instalar openpyxl: python -m pip install openpyxl"
            ) from exc
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
        sheet = workbook.active
        values = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=1, values_only=True), start=1):
            candidate = extract_document_from_row(row, row_index)
            if candidate:
                values.append(candidate)
    else:
        raise ValueError("El archivo debe ser .txt, .xls o .xlsx")

    values = [v for v in values if v]
    if not values:
        raise ValueError("El archivo no contiene ningún RUC válido.")
    return values


def chunk_values(values: List[str], batch_size: int = 100) -> List[List[str]]:
    return [values[i : i + batch_size] for i in range(0, len(values), batch_size)]


def get_output_dir(file_path: str) -> str:
    base_dir = os.path.dirname(file_path) or os.getcwd()
    output_dir = os.path.join(base_dir, "resultados_sunat")
    if os.path.isfile(output_dir):
        output_dir = os.path.join(base_dir, f"resultados_sunat_{random.randint(1000,9999)}")
    try:
        os.makedirs(output_dir, exist_ok=True)
        return output_dir
    except OSError:
        fallback_dir = os.path.join(os.getcwd(), "resultados_sunat")
        os.makedirs(fallback_dir, exist_ok=True)
        return fallback_dir


def clean_val(val: Optional[str], fallback: str = "N/A") -> str:
    if val is None:
        return fallback
    v = parse_cell_to_string(val)
    if not v or v == "-" or v.upper() == "NONE" or v.upper() == "N/A":
        return fallback
    return v


def parse_sunat_zip_results(output_dir: str) -> dict:
    result: dict = {}
    for fname in sorted(os.listdir(output_dir)):
        if not fname.lower().endswith(".zip"):
            continue
        try:
            with zipfile.ZipFile(os.path.join(output_dir, fname), "r") as zf:
                for inner in zf.namelist():
                    if not inner.lower().endswith(".txt") or inner.startswith("__"):
                        continue
                    raw_bytes = zf.read(inner)
                    try:
                        content = raw_bytes.decode("utf-8")
                    except UnicodeDecodeError:
                        content = raw_bytes.decode("latin-1")
                    
                    rows = parse_pipe_delimited_text(content)
                    if not rows:
                        continue
                    start = 1 if rows and "NumeroRuc" in rows[0][0] else 0
                    for row in rows[start:]:
                        if len(row) < 2:
                            continue
                        num_doc = parse_cell_to_string(row[0]).strip("|")
                        if len(num_doc) == 10:
                            num_doc = num_doc.zfill(11)
                        if len(num_doc) != 11 or not num_doc.isdigit():
                            continue
                        nombres   = clean_val(row[1])            if len(row) > 1  else "N/A"
                        estado    = clean_val(row[6])            if len(row) > 6  else "N/A"
                        dep       = clean_val(row[9],  fallback="") if len(row) > 9  else ""
                        prov      = clean_val(row[10], fallback="") if len(row) > 10 else ""
                        dist      = clean_val(row[11], fallback="") if len(row) > 11 else ""
                        direccion = clean_val(row[12])               if len(row) > 12 else "N/A"
                        ubigeo = f"{dep or 'N/A'}/{prov or 'N/A'}/{dist or 'N/A'}" if (dep or prov or dist) else "N/A"
                        result[num_doc] = {"NOMBRES": nombres, "ESTADO": estado, "DIRECCION": direccion, "UBIGEO": ubigeo}
        except Exception as exc:
            print(f"[parse_sunat] error en {fname}: {exc}")
    return result


def _auto_fit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 12)


def _save_wb(wb, path: str) -> str:
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = os.path.splitext(path)[0] + f"_{random.randint(1000,9999)}.xlsx"
        wb.save(alt)
        return alt


def process_date_folder_automatically(parent_folder_path: str, fecha_inicio: str, fecha_fin: str, progress_callback=None) -> tuple:
    f_inicio = datetime.strptime(fecha_inicio, "%Y-%m-%d")
    f_fin = datetime.strptime(fecha_fin, "%Y-%m-%d")
    
    processed_reports = []
    error_items = []

    if not os.path.exists(parent_folder_path):
        raise FileNotFoundError(f"La ruta principal no existe: {parent_folder_path}")

    subdirs = [d for d in os.listdir(parent_folder_path) if os.path.isdir(os.path.join(parent_folder_path, d))]

    for folder_name in sorted(subdirs):
        try:
            folder_date = datetime.strptime(folder_name.strip(), "%Y-%m-%d")
        except ValueError:
            continue

        if f_inicio <= folder_date <= f_fin:
            day_folder_path = os.path.join(parent_folder_path, folder_name)
            
            if progress_callback:
                progress_callback("folder_processing", {"folder": folder_name})

            source_file = None
            for item in os.listdir(day_folder_path):
                item_path = os.path.join(day_folder_path, item)
                item_lower = item.lower()
                if os.path.isfile(item_path) and item_lower.endswith((".xls", ".xlsx", ".txt")):
                    if "resultados" not in item_lower and "actualizado" not in item_lower and item_lower != "ruc.xlsx":
                        source_file = item_path
                        break

            if not source_file:
                err_msg = "No se encontró archivo de datos de entrada"
                error_items.append({"folder": folder_name, "path": day_folder_path, "error": err_msg})
                if progress_callback:
                    progress_callback("folder_error", {"folder": folder_name, "error": err_msg})
                continue

            try:
                input_data = load_input_excel_data(source_file)
                rucs_only = [parse_cell_to_string(r["DNI"]) for r in input_data if parse_cell_to_string(r["DNI"])]

                if not rucs_only:
                    err_msg = "El archivo no contiene RUCs de 11 dígitos válidos"
                    error_items.append({"folder": folder_name, "path": source_file, "error": err_msg})
                    if progress_callback:
                        progress_callback("folder_error", {"folder": folder_name, "error": err_msg})
                    continue

                ruc_wb = Workbook()
                ruc_ws = ruc_wb.active
                ruc_ws.title = "RUCs"
                ruc_ws.append(["DNI"])
                for ruc in rucs_only:
                    ruc_ws.append([ruc])

                ruc_file_target = os.path.join(day_folder_path, "RUC.xlsx")
                ruc_wb.save(ruc_file_target)

                res_c, res_a = process_full_consolidado(ruc_file_target, progress_callback=progress_callback)
                processed_reports.append(res_c)

            except Exception as exc:
                err_msg = str(exc)
                error_items.append({"folder": folder_name, "path": source_file if source_file else day_folder_path, "error": err_msg})
                if progress_callback:
                    progress_callback("folder_error", {"folder": folder_name, "error": err_msg})

    return processed_reports, error_items


def process_full_consolidado(file_path: str, progress_callback=None) -> tuple:
    input_data = load_input_excel_data(file_path)
    ruc_to_query = [normalize_ruc_text(parse_cell_to_string(r["DNI"])) for r in input_data if len(parse_cell_to_string(r["DNI"])) == 11 and parse_cell_to_string(r["DNI"]).isdigit()]
    total = len(ruc_to_query)
    output_dir = get_output_dir(file_path)

    if total == 0:
        sunat_data: dict = {}
        c = build_consolidated_output(input_data, sunat_data, output_dir)
        a = build_updated_input_excel(input_data, sunat_data, file_path)
        return c, a

    batches = chunk_values(ruc_to_query, batch_size=100)
    total_batches = len(batches)

    if progress_callback:
        progress_callback("start", {"total_rucs": total, "total_batches": total_batches})

    session = requests.Session()
    session.headers.update(HEADERS)

    for batch_index, batch in enumerate(batches, start=1):
        start_idx = (batch_index - 1) * 100 + 1
        end_idx = min(batch_index * 100, total)

        if progress_callback:
            progress_callback("batch_start", {
                "batch_index": batch_index, "total_batches": total_batches,
                "start_idx": start_idx, "end_idx": end_idx, "batch_len": len(batch),
            })

        with tempfile.NamedTemporaryFile("w", delete=False, suffix=".txt", encoding="utf-8") as tf:
            tf.writelines(f"{v}\n" for v in batch)
            temp_path = tf.name

        try:
            get_resp = session.get(BASE_URL, timeout=30)
            get_resp.encoding = get_resp.apparent_encoding
            get_resp.raise_for_status()

            token = generate_sunat_token(52)
            payload = {"accion": "consArchivo", "tokenArch": token}
            with open(temp_path, "rb") as archivo:
                files = {"archivo": (os.path.basename(temp_path), archivo, "text/plain")}
                response = session.post(BASE_URL, data=payload, files=files, timeout=120)
            
            response.encoding = response.apparent_encoding
            response.raise_for_status()
            save_sunat_response(response, output_dir, batch_index)
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    sunat_data = parse_sunat_zip_results(output_dir)
    consolidado = build_consolidated_output(input_data, sunat_data, output_dir)
    actualizado = build_updated_input_excel(input_data, sunat_data, file_path)

    return consolidado, actualizado


def consult_ruc_file(file_path: str, progress_callback=None) -> List[tuple[int, str, requests.Response]]:
    values = load_rucs_from_file(file_path)
    total_rucs = len(values)
    batches = chunk_values(values, batch_size=100)
    total_batches = len(batches)
    output_dir = get_output_dir(file_path)
    results: List[tuple[int, str, requests.Response]] = []

    if progress_callback:
        progress_callback("start", {"total_rucs": total_rucs, "total_batches": total_batches})

    session = requests.Session()
    session.headers.update(HEADERS)

    for batch_index, batch in enumerate(batches, start=1):
        start_idx = (batch_index - 1) * 100 + 1
        end_idx = min(batch_index * 100, total_rucs)

        if progress_callback:
            progress_callback("batch_start", {
                "batch_index": batch_index, "total_batches": total_batches,
                "start_idx": start_idx, "end_idx": end_idx, "batch_len": len(batch)
            })

        with tempfile.NamedTemporaryFile("w", delete=False, suffix=".txt", encoding="utf-8") as temp_file:
            temp_file.writelines(f"{v}\n" for v in batch)
            temp_path = temp_file.name

        try:
            response = session.get(BASE_URL, timeout=30)
            response.encoding = response.apparent_encoding
            response.raise_for_status()

            token = generate_sunat_token(52)
            payload = {"accion": "consArchivo", "tokenArch": token}
            with open(temp_path, "rb") as archivo:
                files = {"archivo": (os.path.basename(temp_path), archivo, "text/plain")}
                response = session.post(BASE_URL, data=payload, files=files, timeout=120)

            response.encoding = response.apparent_encoding
            response.raise_for_status()
            saved_path = save_sunat_response(response, output_dir, batch_index)
            results.append((batch_index, saved_path, response))

            if progress_callback:
                if "_ERROR" in saved_path:
                    progress_callback("batch_warning", {
                        "batch_index": batch_index, "total_batches": total_batches,
                        "saved_path": saved_path, "warning_msg": "RUCs no válidos en este lote."
                    })
                else:
                    progress_callback("batch_success", {
                        "batch_index": batch_index, "total_batches": total_batches,
                        "saved_path": saved_path, "status_code": response.status_code
                    })
        except Exception as exc:
            if progress_callback:
                progress_callback("batch_error", {"batch_index": batch_index, "error_msg": str(exc)})
            raise
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    if progress_callback:
        progress_callback("consolidate_start", {"output_dir": output_dir})

    consolidate_results(output_dir)

    if progress_callback:
        progress_callback("consolidate_success", {
            "output_dir": output_dir,
            "consolidated_path": os.path.normpath(os.path.join(output_dir, "resultados_consolidados.xlsx"))
        })

    return results


def consolidate_results(output_dir: str) -> Optional[str]:
    excel_path = os.path.join(output_dir, "resultados_consolidados.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"

    headers = [
        "NumDocumento", "TipoDocumento", "Apellido_Paterno", "Apellido_Materno",
        "Nombres", "Direccion", "EstadoCivil", "Ubigeo", "Restriccion"
    ]
    ws.append(headers)
    has_data = False

    for file_name in sorted(os.listdir(output_dir)):
        if not file_name.lower().endswith(".zip"):
            continue

        zip_path = os.path.join(output_dir, file_name)
        try:
            with zipfile.ZipFile(zip_path, "r") as zf:
                for name in zf.namelist():
                    if name.lower().endswith(".txt") and not name.startswith("__"):
                        raw_bytes = zf.read(name)
                        try:
                            content = raw_bytes.decode("utf-8")
                        except UnicodeDecodeError:
                            content = raw_bytes.decode("latin-1")
                        
                        rows = parse_pipe_delimited_text(content)
                        if not rows:
                            continue

                        start_row_idx = 1 if rows and "NumeroRuc" in rows[0][0] else 0

                        for row in rows[start_row_idx:]:
                            if len(row) < 2:
                                continue

                            num_doc = clean_val(row[0], fallback="") if len(row) > 0 else ""
                            num_doc = parse_cell_to_string(num_doc).strip("|")
                            if len(num_doc) == 10:
                                num_doc = num_doc.zfill(11)
                            if not num_doc:
                                continue

                            nombres = clean_val(row[1], fallback="N/A") if len(row) > 1 else "N/A"
                            estado = clean_val(row[6], fallback="N/A") if len(row) > 6 else "N/A"
                            dep = clean_val(row[9], fallback="") if len(row) > 9 else ""
                            prov = clean_val(row[10], fallback="") if len(row) > 10 else ""
                            dist = clean_val(row[11], fallback="") if len(row) > 11 else ""
                            direccion = clean_val(row[12], fallback="N/A") if len(row) > 12 else "N/A"

                            ubigeo = f"{dep or 'N/A'}/{prov or 'N/A'}/{dist or 'N/A'}" if (dep or prov or dist) else "N/A"

                            new_row = [
                                num_doc, "RUC", "N/A", "N/A",
                                nombres, direccion, "N/A", ubigeo, estado
                            ]
                            ws.append(new_row)
                            has_data = True
        except Exception as e:
            print(f"Error consolidando archivo {file_name}: {e}")

    if has_data:
        _auto_fit(ws)
        return _save_wb(wb, excel_path)
    return None


def build_consolidated_output(input_data: List[dict], sunat_data: dict, output_dir: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    ws.append(["NumDocumento", "TipoDocumento", "Apellido_Paterno", "Apellido_Materno",
               "Nombres", "Direccion", "EstadoCivil", "Ubigeo", "Restriccion"])

    for r in input_data:
        dni_val = parse_cell_to_string(r["DNI"])
        if len(dni_val) == 11 and dni_val.isdigit():
            s = sunat_data.get(dni_val, {})
            ws.append([dni_val, "RUC", "N/A", "N/A",
                       s.get("NOMBRES", "N/A"), s.get("DIRECCION", "N/A"),
                       "N/A", s.get("UBIGEO", "N/A"), s.get("ESTADO", "N/A")])
        else:
            ws.append([dni_val, "DNI", r["APELLIDO_PATERNO"], r["APELLIDO_MATERNO"],
                       r["NOMBRES"], r["DIRECCION"], r["ESTADOCIVIL"], r["UBIGEO"], r["RESTRICCION"]])

    _auto_fit(ws)
    return _save_wb(wb, os.path.join(output_dir, "resultados_consolidados.xlsx"))


def build_updated_input_excel(input_data: List[dict], sunat_data: dict, source_path: str) -> str:
    base_dir = os.path.dirname(source_path)
    base_name = os.path.splitext(os.path.basename(source_path))[0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Actualizado"
    ws.append(["DNI", "Apellido_Paterno", "Apellido_Materno", "Nombres",
               "Direccion", "EstadoCivil", "Ubigeo", "Restriccion"])

    for r in input_data:
        dni_val = parse_cell_to_string(r["DNI"])
        if len(dni_val) == 11 and dni_val.isdigit():
            s = sunat_data.get(dni_val, {})
            ws.append([dni_val, "N/A", "N/A",
                       s.get("NOMBRES", "N/A"),
                       s.get("DIRECCION", "N/A"), "N/A",
                       s.get("UBIGEO", "N/A"), s.get("ESTADO", "N/A")])
        else:
            ws.append([dni_val, r["APELLIDO_PATERNO"], r["APELLIDO_MATERNO"], r["NOMBRES"],
                       r["DIRECCION"], r["ESTADOCIVIL"], r["UBIGEO"], r["RESTRICCION"]])

    _auto_fit(ws)
    return _save_wb(wb, os.path.join(base_dir, f"{base_name}_actualizado.xlsx"))


_COL_ALIAS: dict = {
    "DNI":              ["DNI", "NUMDNI", "NUM_DNI", "NUMDOCUMENTO", "NUM_DOCUMENTO", "DOCUMENTO", "RUC"],
    "APELLIDO_PATERNO": ["APELLIDO_PATERNO", "APELLIDOPATERNO", "APE_PATERNO", "APEPATERNO"],
    "APELLIDO_MATERNO": ["APELLIDO_MATERNO", "APELLIDOMATERNO", "APE_MATERNO", "APEMATERNO"],
    "NOMBRES":          ["NOMBRES", "NOMBRE", "NAMES", "NAME"],
    "ESTADOCIVIL":      ["ESTADOCIVIL", "ESTADOCIVIL", "ESTADO_CIVIL"],
    "SEXO":             ["SEXO", "GENERO", "GENDER"],
    "DIRECCION":        ["DIRECCION", "DIRECCION", "ADDRESS"],
    "UBIGEO":           ["UBIGEO", "COD_UBIGEO", "CODUBIGEO"],
    "RESTRICCION":      ["RESTRICCION", "RESTRICCION"],
}


def _norm_col(s: str) -> str:
    s = s.upper().strip()
    for src, dst in [("Á","A"),("É","E"),("Í","I"),("Ó","O"),("Ú","U"),("Ñ","N"),
                     ("á","A"),("é","E"),("í","I"),("ó","O"),("ú","U"),("ñ","N")]:
        s = s.replace(src, dst)
    return re.sub(r"[\s_\-]", "", s)


def _find_col(header: list, field: str) -> int:
    aliases = [_norm_col(a) for a in _COL_ALIAS.get(field, [field])]
    for i, h in enumerate(header):
        if _norm_col(str(h)) in aliases:
            return i
    return -1


def load_input_excel_data(file_path: str) -> List[dict]:
    try:
        from openpyxl import load_workbook as _lw
    except ImportError as exc:
        raise ImportError("Instale openpyxl: python -m pip install openpyxl") from exc

    wb = _lw(filename=file_path, read_only=True, data_only=True)
    ws = wb.active

    col_map: dict = {}
    rows_data: List[dict] = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [parse_cell_to_string(c) for c in row]

        if row_idx == 1:
            for field in _COL_ALIAS:
                idx = _find_col(cells, field)
                if idx >= 0:
                    col_map[field] = idx
            continue

        if not any(cells):
            continue

        _NULL_VALS = {"null", "none", "n/a", "-", ""}

        dni_idx = col_map.get("DNI", -1)
        dni_raw = cells[dni_idx] if 0 <= dni_idx < len(cells) else ""
        dni_clean = re.sub(r"[^0-9]", "", dni_raw)
        
        if len(dni_clean) == 10:
            dni_clean = dni_clean.zfill(11)

        if not dni_clean or len(dni_clean) != 11:
            continue

        def _get(field: str, _cells=cells) -> str:
            idx = col_map.get(field, -1)
            if 0 <= idx < len(_cells):
                v = _cells[idx].strip()
                return "N/A" if v.lower() in _NULL_VALS else v
            return "N/A"

        rows_data.append({
            "DNI":               dni_clean,
            "APELLIDO_PATERNO": _get("APELLIDO_PATERNO"),
            "APELLIDO_MATERNO": _get("APELLIDO_MATERNO"),
            "NOMBRES":          _get("NOMBRES"),
            "ESTADOCIVIL":      _get("ESTADOCIVIL"),
            "SEXO":             _get("SEXO"),
            "DIRECCION":        _get("DIRECCION"),
            "UBIGEO":           _get("UBIGEO"),
            "RESTRICCION":      _get("RESTRICCION"),
        })

    if not rows_data:
        raise ValueError(f"No se encontraron filas con datos válidos en el archivo: {os.path.basename(file_path)}")
    return rows_data


def run_gui() -> None:
    if tk is None:
        print("No se puede ejecutar la GUI porque Tkinter no está disponible.")
        return

    C_BG      = "#0F172A"
    C_SURFACE = "#1E293B"
    C_BORDER  = "#334155"
    C_ACCENT  = "#0EA5E9"
    C_SUCCESS = "#10B981"
    C_WARNING = "#F59E0B"
    C_ERROR   = "#EF4444"
    C_TEXT    = "#F1F5F9"
    C_MUTED   = "#64748B"
    C_INPUT   = "#0B1120"
    F         = "Segoe UI"

    root = tk.Tk()
    root.title("SUNAT RUC Pro - Sistema Integral")
    root.geometry("920x740")
    root.configure(bg=C_BG)
    root.minsize(760, 600)

    style = ttk.Style()
    style.theme_use('default')
    style.configure('TNotebook', background=C_BG, borderwidth=0)
    style.configure('TNotebook.Tab', background=C_SURFACE, foreground=C_TEXT, padding=[16, 8], font=(F, 9, 'bold'))
    style.map('TNotebook.Tab', background=[('selected', C_ACCENT)], foreground=[('selected', '#0F172A')])

    notebook = ttk.Notebook(root)
    notebook.pack(fill=tk.BOTH, expand=True, padx=16, pady=16)

    # ==================== PESTAÑA 1: MODO MANUAL INTEGRADO ====================
    tab1 = tk.Frame(notebook, bg=C_BG)
    notebook.add(tab1, text="  📁 Modo Manual (Crear RUC y Enviar)  ")

    pad1 = tk.Frame(tab1, bg=C_BG)
    pad1.pack(fill=tk.BOTH, expand=True, padx=16, pady=16)

    tk.Label(pad1, text="SUNAT RUC Pro", font=(F, 18, "bold"), bg=C_BG, fg=C_TEXT).pack(anchor=tk.W)
    tk.Label(pad1, text="Genera el archivo RUC automáticamente y consulta masivo en SUNAT", font=(F, 9), bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(2, 12))
    tk.Frame(pad1, height=1, bg=C_BORDER).pack(fill=tk.X, pady=(0, 12))

    tk.Label(pad1, text="ARCHIVO DE ENTRADA (DIARIO)", font=(F, 8, "bold"), bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(0, 4))
    
    selected_path_manual = tk.StringVar()
    input_wrap1 = tk.Frame(pad1, bg=C_INPUT, padx=12, pady=8)
    input_wrap1.pack(fill=tk.X, pady=(0, 8))
    entry_manual = tk.Entry(input_wrap1, textvariable=selected_path_manual, font=(F, 10), bg=C_INPUT, fg=C_TEXT, insertbackground=C_TEXT, relief=tk.FLAT)
    entry_manual.pack(fill=tk.X)

    btns1 = tk.Frame(pad1, bg=C_BG)
    btns1.pack(fill=tk.X, pady=(0, 12))

    def select_file_manual(initial_dir=None):
        target_dir = os.path.abspath(initial_dir) if initial_dir and os.path.isdir(initial_dir) else os.getcwd()
        path = filedialog.askopenfilename(
            title="Selecciona el archivo de la fecha",
            initialdir=target_dir,
            filetypes=[("Archivos de Excel y Texto", "*.txt *.xlsx *.xls"), ("Todos los archivos", "*.*")]
        )
        if path:
            selected_path_manual.set(path)
            _log1(f"Archivo seleccionado: {path}\n", "success")

    btn_select1 = tk.Button(btns1, text="Seleccionar archivo", font=(F, 9), bg=C_SURFACE, fg=C_TEXT, relief=tk.FLAT, padx=12, pady=6, command=lambda: select_file_manual(None), cursor="hand2")
    btn_select1.pack(side=tk.LEFT)

    def procesar_y_enviar_manual(path_override=None):
        path = path_override or selected_path_manual.get()
        if not path:
            messagebox.showwarning("Sin archivo", "Seleccione un archivo de entrada primero.")
            return

        btn_select1.config(state=tk.DISABLED)
        btn_send1.config(state=tk.DISABLED)

        def _worker():
            def cb(evt, data):
                if evt == "start":
                    _log1(f"   ↳ {data['total_rucs']} RUCs encontrados — {data['total_batches']} lotes\n")
                elif evt == "batch_start":
                    _log1(f"    → Consultando lotes del {data['start_idx']} al {data['end_idx']}...\n")

            try:
                _log1(f"\n[1/3] Leyendo archivo y extrayendo RUCs limpios...\n")
                input_data = load_input_excel_data(path)
                rucs_only = [parse_cell_to_string(r["DNI"]) for r in input_data if parse_cell_to_string(r["DNI"])]

                if not rucs_only:
                    raise ValueError("El archivo no contiene RUCs de 11 dígitos válidos.")

                base_dir = os.path.dirname(path)
                ruc_file_target = os.path.join(base_dir, "RUC.xlsx")

                ruc_wb = Workbook()
                ruc_ws = ruc_wb.active
                ruc_ws.title = "RUCs"
                ruc_ws.append(["RUC"])
                for ruc in rucs_only:
                    ruc_ws.append([ruc])
                ruc_wb.save(ruc_file_target)

                _log1(f"[2/3] Archivo RUC creado con éxito en: {ruc_file_target}\n", "success")
                _log1(f"[3/3] Enviando a SUNAT y generando consolidados...\n")

                consolidado, actualizado = process_full_consolidado(ruc_file_target, progress_callback=cb)

                _log1(f"  ✓ Carpeta de salida creada: {os.path.dirname(consolidado)}\n", "success")
                _log1(f"  ✓ Archivo Consolidado: {consolidado}\n", "success")
                _log1(f"  ✓ Archivo Actualizado: {actualizado}\n", "success")
                
                messagebox.showinfo("Completado", "¡Proceso finalizado con éxito!\nRevisa la carpeta 'resultados_sunat'.")
            except Exception as e:
                _log1(f"\nError: {e}\n", "error")
                messagebox.showerror("Error", str(e))
            finally:
                btn_select1.config(state=tk.NORMAL)
                btn_send1.config(state=tk.NORMAL)

        threading.Thread(target=_worker, daemon=True).start()

    btn_send1 = tk.Button(btns1, text="Crear RUC y Enviar a SUNAT", font=(F, 9, "bold"), bg=C_SUCCESS, fg="#0F172A", relief=tk.FLAT, padx=16, pady=6, command=lambda: procesar_y_enviar_manual(None), cursor="hand2")
    btn_send1.pack(side=tk.LEFT, padx=(8, 0))

    tk.Label(pad1, text="BITÁCORA", font=(F, 8, "bold"), bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(0, 4))
    log1 = scrolledtext.ScrolledText(pad1, font=("Consolas", 9), bg=C_INPUT, fg="#94A3B8", relief=tk.FLAT, padx=12, pady=8)
    log1.pack(fill=tk.BOTH, expand=True)
    log1.tag_config("success", foreground=C_SUCCESS)
    log1.tag_config("error", foreground=C_ERROR)

    def _log1(msg, tag=""):
        log1.config(state=tk.NORMAL)
        log1.insert(tk.END, msg, tag)
        log1.see(tk.END)
        log1.config(state=tk.DISABLED)

    _log1("Consola manual lista. Seleccione su archivo diario para iniciar.\n")

    # ==================== PESTAÑA 2: AUTOMATIZACIÓN POR FECHAS ====================
    tab2 = tk.Frame(notebook, bg=C_BG)
    notebook.add(tab2, text="  ⚙️ Automatización Masiva (DATOS SAT)  ")

    pad2 = tk.Frame(tab2, bg=C_BG)
    pad2.pack(fill=tk.BOTH, expand=True, padx=16, pady=16)

    tk.Label(pad2, text="Automatización Diaria DATOS SAT", font=(F, 18, "bold"), bg=C_BG, fg=C_TEXT).pack(anchor=tk.W)
    tk.Label(pad2, text="Generación automática de RUC y consulta masiva por rangos de fecha", font=(F, 9), bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(2, 12))
    tk.Frame(pad2, height=1, bg=C_BORDER).pack(fill=tk.X, pady=(0, 12))

    tk.Label(pad2, text="CARPETA PRINCIPAL (DATOS SAT)", font=(F, 8, "bold"), bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(0, 4))
    
    selected_folder_auto = tk.StringVar()
    f_wrap2 = tk.Frame(pad2, bg=C_BG)
    f_wrap2.pack(fill=tk.X, pady=(0, 10))

    entry_folder = tk.Entry(f_wrap2, textvariable=selected_folder_auto, font=(F, 10), bg=C_INPUT, fg=C_TEXT, insertbackground=C_TEXT, relief=tk.FLAT)
    entry_folder.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=6, padx=(0, 8))
    
    def select_parent_folder():
        path = filedialog.askdirectory(title="Seleccionar carpeta DATOS SAT")
        if path:
            selected_folder_auto.set(path)
            _log2(f"Carpeta principal: {path}\n")

    btn_folder2 = tk.Button(f_wrap2, text="Elegir Carpeta", font=(F, 9), bg=C_SURFACE, fg=C_TEXT, relief=tk.FLAT, padx=12, command=select_parent_folder, cursor="hand2")
    btn_folder2.pack(side=tk.RIGHT)

    dates_frame2 = tk.Frame(pad2, bg=C_BG)
    dates_frame2.pack(fill=tk.X, pady=(0, 10))

    years = [str(y) for y in range(2024, 2031)]
    months = [f"{m:02d}" for m in range(1, 13)]
    days = [f"{d:02d}" for d in range(1, 32)]

    # Fecha Desde
    tk.Label(dates_frame2, text="Desde:", font=(F, 9), bg=C_BG, fg=C_TEXT).pack(side=tk.LEFT)
    anio_ini = ttk.Combobox(dates_frame2, values=years, width=5, state="readonly", font=(F, 9))
    anio_ini.set("2026")
    anio_ini.pack(side=tk.LEFT, padx=(4, 2))
    tk.Label(dates_frame2, text="-", bg=C_BG, fg=C_TEXT).pack(side=tk.LEFT)

    mes_ini = ttk.Combobox(dates_frame2, values=months, width=3, state="readonly", font=(F, 9))
    mes_ini.set("07")
    mes_ini.pack(side=tk.LEFT, padx=2)
    tk.Label(dates_frame2, text="-", bg=C_BG, fg=C_TEXT).pack(side=tk.LEFT)

    dia_ini = ttk.Combobox(dates_frame2, values=days, width=3, state="readonly", font=(F, 9))
    dia_ini.set("22")
    dia_ini.pack(side=tk.LEFT, padx=(2, 16))

    # Fecha Hasta
    tk.Label(dates_frame2, text="Hasta:", font=(F, 9), bg=C_BG, fg=C_TEXT).pack(side=tk.LEFT)
    anio_fin = ttk.Combobox(dates_frame2, values=years, width=5, state="readonly", font=(F, 9))
    anio_fin.set("2026")
    anio_fin.pack(side=tk.LEFT, padx=(4, 2))
    tk.Label(dates_frame2, text="-", bg=C_BG, fg=C_TEXT).pack(side=tk.LEFT)

    mes_fin = ttk.Combobox(dates_frame2, values=months, width=3, state="readonly", font=(F, 9))
    mes_fin.set("07")
    mes_fin.pack(side=tk.LEFT, padx=2)
    tk.Label(dates_frame2, text="-", bg=C_BG, fg=C_TEXT).pack(side=tk.LEFT)

    dia_fin = ttk.Combobox(dates_frame2, values=days, width=3, state="readonly", font=(F, 9))
    dia_fin.set("23")
    dia_fin.pack(side=tk.LEFT, padx=(2, 16))

    # PANEL INFERIOR PARA ERRORES PENDIENTES
    err_panel = tk.Frame(pad2, bg=C_SURFACE, padx=10, pady=8)
    err_panel.pack(fill=tk.BOTH, expand=True, pady=(4, 8))

    tk.Label(err_panel, text="⚠️ CARPETAS CON ERROR / PENDIENTES DE REVISIÓN MANUAL:", font=(F, 8, "bold"), bg=C_SURFACE, fg=C_WARNING).pack(anchor=tk.W, pady=(0, 4))
    
    err_list_frame = tk.Frame(err_panel, bg=C_SURFACE)
    err_list_frame.pack(fill=tk.BOTH, expand=True)

    err_scrollbar = tk.Scrollbar(err_list_frame, orient=tk.VERTICAL)
    err_listbox = tk.Listbox(err_list_frame, font=("Consolas", 9), bg=C_INPUT, fg=C_TEXT, selectbackground=C_ACCENT, selectforeground="#0F172A", relief=tk.FLAT, yscrollcommand=err_scrollbar.set, height=5)
    err_scrollbar.config(command=err_listbox.yview)
    
    err_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    err_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    pending_errors_dict = {}

    err_btns_frame = tk.Frame(err_panel, bg=C_SURFACE)
    err_btns_frame.pack(fill=tk.X, pady=(6, 0))

    def resolver_seleccionado_manual():
        selection = err_listbox.curselection()
        if not selection:
            messagebox.showwarning("Sin selección", "Seleccione una carpeta con error de la lista para resolverla.")
            return
        idx = selection[0]
        item_text = err_listbox.get(idx)
        folder_key = item_text.split(" | ")[0].replace("📅 ", "").strip()
        
        if folder_key in pending_errors_dict:
            target_folder_path = pending_errors_dict[folder_key]
            
            notebook.select(0)
            
            if os.path.exists(target_folder_path):
                select_file_manual(initial_dir=target_folder_path)
            else:
                parent_dir = os.path.dirname(target_folder_path)
                select_file_manual(initial_dir=parent_dir)
            
            err_listbox.delete(idx)
            del pending_errors_dict[folder_key]
            _log2(f"  ✓ Carpeta de fecha '{folder_key}' abierta en explorador y removida de pendientes.\n", "success")

    def quitar_de_lista():
        selection = err_listbox.curselection()
        if not selection:
            return
        idx = selection[0]
        item_text = err_listbox.get(idx)
        folder_key = item_text.split(" | ")[0].replace("📅 ", "").strip()
        
        err_listbox.delete(idx)
        if folder_key in pending_errors_dict:
            del pending_errors_dict[folder_key]
        _log2(f"  ✓ Eliminado de pendientes: {folder_key}\n", "success")

    btn_resolver = tk.Button(err_btns_frame, text="🛠️ Resolver Manualmente", font=(F, 8, "bold"), bg=C_ACCENT, fg="#0F172A", relief=tk.FLAT, padx=10, pady=4, command=resolver_seleccionado_manual, cursor="hand2")
    btn_resolver.pack(side=tk.LEFT)

    btn_quitar = tk.Button(err_btns_frame, text="🗑️ Quitar de la lista", font=(F, 8), bg=C_SURFACE, fg=C_TEXT, relief=tk.FLAT, padx=10, pady=4, command=quitar_de_lista, cursor="hand2")
    btn_quitar.pack(side=tk.LEFT, padx=(8, 0))

    def ejecutar_automatizacion():
        folder_path = selected_folder_auto.get()
        f_ini = f"{anio_ini.get()}-{mes_ini.get()}-{dia_ini.get()}"
        f_fin = f"{anio_fin.get()}-{mes_fin.get()}-{dia_fin.get()}"

        if not folder_path:
            messagebox.showwarning("Sin carpeta", "Seleccione la carpeta principal 'DATOS SAT'.")
            return

        err_listbox.delete(0, tk.END)
        pending_errors_dict.clear()

        btn_folder2.config(state=tk.DISABLED)
        btn_run2.config(state=tk.DISABLED)

        def _worker():
            def cb(evt, data):
                if evt == "folder_processing":
                    _log2(f"\n📂 Procesando fecha: {data['folder']}\n", "success")
                elif evt == "start":
                    _log2(f"   ↳ {data['total_rucs']} RUCs encontrados\n")
                elif evt == "folder_error":
                    _log2(f"   ❌ Error: {data['error']}\n", "error")

            try:
                _log2(f"\n[INICIO] Barrido automático del {f_ini} al {f_fin}...\n", "success")
                reports, error_items = process_date_folder_automatically(folder_path, f_ini, f_fin, progress_callback=cb)
                
                for err in error_items:
                    f_name = err["folder"]
                    f_path = err["path"]
                    f_err = err["error"]
                    real_folder_path = f_path if os.path.isdir(f_path) else os.path.dirname(f_path)
                    pending_errors_dict[f_name] = real_folder_path
                    root.after(0, lambda fn=f_name, fe=f_err: err_listbox.insert(tk.END, f"📅 {fn} | Error: {fe}"))

                _log2(f"\n[ÉXITO] Procesadas: {len(reports)} carpetas | Errores pendientes: {len(error_items)}\n", "success")
                messagebox.showinfo("Completado", f"Automatización finalizada.\nCarpetas con error enviadas al panel inferior: {len(error_items)}")
            except Exception as exc:
                _log2(f"\nError general: {exc}\n", "error")
                messagebox.showerror("Error", str(exc))
            finally:
                root.after(0, lambda: [
                    btn_folder2.config(state=tk.NORMAL),
                    btn_run2.config(state=tk.NORMAL)
                ])

        threading.Thread(target=_worker, daemon=True).start()

    btn_run2 = tk.Button(dates_frame2, text="Iniciar Automatización", font=(F, 9, "bold"), bg=C_SUCCESS, fg="#0F172A", relief=tk.FLAT, padx=16, pady=6, command=ejecutar_automatizacion, cursor="hand2")
    btn_run2.pack(side=tk.RIGHT)

    tk.Label(pad2, text="BITÁCORA DE AUTOMATIZACIÓN", font=(F, 8, "bold"), bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(0, 4))
    log2 = scrolledtext.ScrolledText(pad2, font=("Consolas", 9), bg=C_INPUT, fg="#94A3B8", relief=tk.FLAT, padx=12, pady=8, height=6)
    log2.pack(fill=tk.BOTH, expand=True)
    log2.tag_config("success", foreground=C_SUCCESS)
    log2.tag_config("error", foreground=C_ERROR)

    def _log2(msg, tag=""):
        log2.config(state=tk.NORMAL)
        log2.insert(tk.END, msg, tag)
        log2.see(tk.END)
        log2.config(state=tk.DISABLED)

    _log2("Módulo automático listo. Rutas absolutas forzadas para el explorador.\n")

    root.mainloop()


def main() -> int:
    run_gui()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())