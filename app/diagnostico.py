"""Diagnóstico de un archivo de entrada para V10.py.

Uso:
    python diagnostico.py "ruta\\DNI 20072026 - Respuesta.xlsx"

Muestra, por cada hoja: dimensiones, primeras 5 filas (con tipo de dato de cada
celda), columnas detectadas con _COL_ALIAS en las primeras 10 filas y el conteo de
longitudes (solo dígitos) de los números de documento.
No modifica nada ni se conecta a internet.
"""
import os
import re
import sys
from collections import Counter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # consola Windows cp1252

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from V10 import _COL_ALIAS, _norm_col, parse_cell_to_string  # noqa: E402

MAX_HEADER_SCAN = 10


def read_sheets(path: str) -> dict:
    """Devuelve {nombre_hoja: [filas como tuplas]}."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        print(f"Hoja activa (ws.active): {wb.active.title!r}")
        return {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}
    if ext == ".xls":
        try:
            import xlrd
        except ImportError:
            sys.exit("Para .xls instale xlrd: python -m pip install xlrd")
        book = xlrd.open_workbook(path)
        return {sh.name: [tuple(sh.row_values(i)) for i in range(sh.nrows)] for sh in book.sheets()}
    if ext == ".txt":
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return {"(txt)": [(line.rstrip("\r\n"),) for line in f]}
    sys.exit(f"Extensión no soportada: {ext}")


def detect_columns(row) -> dict:
    """Aplica la misma lógica que load_input_excel_data (fila -> {campo: índice})."""
    cells = [parse_cell_to_string(c) for c in row]
    found = {}
    for field, aliases in _COL_ALIAS.items():
        norm_aliases = {_norm_col(a) for a in aliases}
        for i, h in enumerate(cells):
            if _norm_col(str(h)) in norm_aliases:
                found[field] = i
                break
    return found


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    path = sys.argv[1]
    if not os.path.isfile(path):
        sys.exit(f"No existe el archivo: {path}")

    print(f"Archivo: {path}")
    print(f"Tamaño: {os.path.getsize(path)} bytes")
    sheets = read_sheets(path)
    print(f"Hojas: {list(sheets.keys())}\n")

    for name, rows in sheets.items():
        print("=" * 70)
        n_cols = max((len(r) for r in rows), default=0)
        print(f"HOJA {name!r}: {len(rows)} filas x {n_cols} columnas")

        print("\n-- Primeras 5 filas (valor repr : tipo) --")
        for i, row in enumerate(rows[:5], start=1):
            shown = [f"{c!r}:{type(c).__name__}" for c in row]
            print(f"  fila {i}: {shown}")

        print(f"\n-- Detección de encabezado con _COL_ALIAS (primeras {MAX_HEADER_SCAN} filas) --")
        header_row = None
        for i, row in enumerate(rows[:MAX_HEADER_SCAN], start=1):
            found = detect_columns(row)
            if found:
                print(f"  fila {i}: {found}")
                if header_row is None and "DNI" in found:
                    header_row = (i, found)
        if header_row is None:
            print("  ⚠ Ninguna fila matchea la columna de documento (DNI).")
            print("  Encabezados normalizados de las primeras 3 filas:")
            for i, row in enumerate(rows[:3], start=1):
                print(f"    fila {i}: {[_norm_col(parse_cell_to_string(c)) for c in row]}")
        else:
            print(f"  → Encabezado en fila {header_row[0]}; V10 actual SOLO mira la fila 1.")

        print("\n-- Conteo de longitudes de documento (solo dígitos) --")
        if header_row is not None:
            start, col = header_row[0], header_row[1]["DNI"]
            values = [r[col] if col < len(r) else None for r in rows[start:]]
            print(f"  (columna índice {col}, desde la fila {start + 1})")
        else:
            col = 0
            values = [r[0] if r else None for r in rows]
            print("  (sin encabezado detectado: se usa la primera columna)")

        lengths = Counter()
        types = Counter()
        examples = {}
        for v in values:
            if v is None or str(v).strip() == "":
                lengths["vacío"] += 1
                continue
            types[type(v).__name__] += 1
            digits = re.sub(r"[^0-9]", "", parse_cell_to_string(v))
            lengths[len(digits)] += 1
            examples.setdefault(len(digits), repr(v))
        for k, cnt in sorted(lengths.items(), key=lambda kv: str(kv[0])):
            ej = f"  ej: {examples[k]}" if k in examples else ""
            print(f"  long {k}: {cnt}{ej}")
        print(f"  tipos de celda: {dict(types)}")
        print()

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
