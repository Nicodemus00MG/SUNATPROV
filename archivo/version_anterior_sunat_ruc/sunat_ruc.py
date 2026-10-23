import argparse
import os
import random
import re
import shutil
import string
import subprocess
import sys
import tempfile
import threading
import zipfile
from html.parser import HTMLParser
from typing import List, Optional
from urllib.parse import unquote, urljoin, urlparse

import requests
from openpyxl import Workbook

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext
except ImportError:
    tk = None

BASE_URL = "https://e-consultaruc.sunat.gob.pe/cl-ti-itmrconsmulruc/jrmS00Alias"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "es-PE,es;q=0.9,en;q=0.8",
    "Referer": BASE_URL,
}


def generate_sunat_token(length: int = 52) -> str:
    alphabet = string.ascii_lowercase + string.digits
    token = ""
    while len(token) < length:
        token += ''.join(random.choice(alphabet) for _ in range(16))
    return token[:length]


def normalize_ruc_text(value: str) -> str:
    value = value.strip()
    if not value:
        return ""
    return value if value.endswith("|") else value + "|"


def extract_document_from_row(row: tuple, row_index: int) -> Optional[str]:
    if not row:
        return None

    cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
    if not cells:
        return None

    if row_index == 1:
        header_text = " ".join(cells).lower()
        if any(keyword in header_text for keyword in ["num", "documento", "tipo", "nombre", "apellido", "ruc", "dni"]):
            return None

    for cell in cells:
        cleaned = re.sub(r"[^0-9]", "", cell)
        if 8 <= len(cleaned) <= 11:
            return normalize_ruc_text(cleaned)

    return None


class SunatHtmlTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: List[List[List[str]]] = []
        self.current_table: List[List[str]] = []
        self.current_row: List[str] = []
        self.current_cell: List[str] = []
        self.in_table = False
        self.in_cell = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        if tag == "table":
            self.in_table = True
            self.current_table = []
        elif self.in_table and tag == "tr":
            self.current_row = []
        elif self.in_table and tag in ("td", "th"):
            self.in_cell = True
            self.current_cell = []

    def handle_endtag(self, tag: str) -> None:
        if self.in_table and tag in ("td", "th"):
            self.in_cell = False
            cell_text = "".join(self.current_cell).strip()
            self.current_row.append(cell_text)
        elif self.in_table and tag == "tr":
            if self.current_row:
                self.current_table.append(self.current_row)
        elif tag == "table" and self.in_table:
            self.in_table = False
            if self.current_table:
                self.tables.append(self.current_table)

    def handle_data(self, data: str) -> None:
        if self.in_cell:
            self.current_cell.append(data)


def parse_html_table(html: str) -> List[List[str]]:
    parser = SunatHtmlTableParser()
    parser.feed(html)
    parser.close()
    return parser.tables[0] if parser.tables else []


def parse_pipe_delimited_text(text: str) -> List[List[str]]:
    rows: List[List[str]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip("\r")
        if not line:
            continue
        rows.append(line.split("|"))
    return rows


def get_fallback_output_dir() -> str:
    fallback_dir = os.path.join(os.getcwd(), "resultados_sunat")
    os.makedirs(fallback_dir, exist_ok=True)
    return fallback_dir


def get_filename_from_response(response: requests.Response) -> Optional[str]:
    header = response.headers.get("content-disposition", "")
    if not header:
        return None
    filename_match = re.search(r"filename\*=UTF-8''(?P<name>[^;]+)", header)
    if filename_match:
        return unquote(filename_match.group("name"))
    filename_match = re.search(r"filename=(?P<name>\"?[^\";]+\"?)", header)
    if filename_match:
        return filename_match.group("name").strip('"')
    return None


def is_archive_content(response: requests.Response) -> bool:
    content_type = response.headers.get("content-type", "").lower()
    if "rar" in content_type or "zip" in content_type:
        return True
    content_start = response.content[:8]
    return content_start.startswith(b"Rar!") or content_start.startswith(b"PK\x03\x04")


def save_binary_response(response: requests.Response, save_path: str) -> str:
    save_path = os.path.normpath(save_path)
    with open(save_path, "wb") as f:
        f.write(response.content)
    return save_path


def extract_archive(archive_path: str, extract_dir: str) -> List[str]:
    os.makedirs(extract_dir, exist_ok=True)

    if zipfile.is_zipfile(archive_path):
        with zipfile.ZipFile(archive_path, "r") as archive:
            archive.extractall(extract_dir)
    elif shutil.which("7z"):
        subprocess.run(["7z", "x", "-y", f"-o{extract_dir}", archive_path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    elif shutil.which("unrar"):
        subprocess.run(["unrar", "x", "-y", archive_path, extract_dir], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    else:
        raise RuntimeError(
            "No se encontró extractor de archivos. Instale 7z o unrar para extraer archivos RAR/ZIP."
        )

    extracted_files: List[str] = []
    for root, _, files in os.walk(extract_dir):
        for file in files:
            extracted_files.append(os.path.join(root, file))
    return extracted_files


def find_archive_link(html: str) -> Optional[str]:
    # 1. Buscar enlace cuyo atributo href termine directamente en .zip o .rar
    matches = re.findall(r'href=["\']([^"\']+\.(?:zip|rar)(?:\?[^"\']*)?)["\']', html, flags=re.IGNORECASE)
    if matches:
        return matches[0]
    
    # 2. Buscar etiqueta <a> cuyo contenido/texto de anclaje contenga .zip o .rar
    # por ejemplo: <a href="url_de_descarga">archivo.zip</a>
    matches_text = re.findall(r'<a\s+[^>]*href=["\']([^"\']+)["\'][^>]*>(?:[^<]*\.(?:zip|rar))</a>', html, flags=re.IGNORECASE)
    if matches_text:
        return matches_text[0]
        
    # 3. Buscar cualquier enlace que tenga patrones de descarga del framework de SUNAT como recurso de respaldo
    matches_download = re.findall(r'href=["\']([^"\']*(?:descargaArchivo|download|descarga)[^"\']*)["\']', html, flags=re.IGNORECASE)
    if matches_download:
        return matches_download[0]

    return None


def get_remote_filename_from_url(url: str) -> str:
    parsed = urlparse(url)
    name = os.path.basename(parsed.path)
    if not name:
        name = "resultado_sunat"
    return name


def save_response_to_excel(text: str, save_path: str) -> str:
    normalized = text.strip()
    if normalized.lower().startswith("<html") or normalized.lower().startswith("<!doctype html"):
        rows = parse_html_table(normalized)
    elif "|" in normalized:
        rows = parse_pipe_delimited_text(normalized)
    else:
        rows = [[normalized]]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SUNAT"

    if not rows:
        sheet.append([normalized])
    else:
        for row in rows:
            sheet.append(row)

    save_path = os.path.normpath(save_path)
    try:
        workbook.save(save_path)
        return save_path
    except PermissionError:
        fallback_dir = get_fallback_output_dir()
        fallback_path = os.path.join(fallback_dir, os.path.basename(save_path))
        workbook.save(fallback_path)
        return fallback_path
    except OSError as exc:
        if getattr(exc, "errno", None) == 13:
            fallback_dir = get_fallback_output_dir()
            fallback_path = os.path.join(fallback_dir, os.path.basename(save_path))
            workbook.save(fallback_path)
            return fallback_path
        raise


def download_archive_from_html(response: requests.Response, output_dir: str) -> Optional[str]:
    if not response.text:
        return None
    archive_link = find_archive_link(response.text)
    if not archive_link:
        return None

    archive_url = urljoin(response.url, archive_link)
    archive_response = requests.get(archive_url, headers=HEADERS, timeout=120)
    archive_response.raise_for_status()

    filename = get_filename_from_response(archive_response) or get_remote_filename_from_url(archive_url)
    _, ext = os.path.splitext(filename)
    if not ext:
        ext = ".zip" if is_archive_content(archive_response) else ".txt"
        filename += ext
    save_path = os.path.join(output_dir, filename)
    return save_binary_response(archive_response, save_path)


def save_sunat_response(response: requests.Response, output_dir: str, batch_index: int) -> str:
    # 1. Comprobar si SUNAT ha retornado una página web con un mensaje de error explícito
    if response.text and ("cuerpo" in response.text or "La aplicación ha retornado" in response.text):
        if "no son válidos o no existen" in response.text or "no son validos o no existen" in response.text:
            filename = f"resultado_sunat_batch_{batch_index}_ERROR_rucs_invalidos.txt"
            save_path = os.path.normpath(os.path.join(output_dir, filename))
            with open(save_path, "w", encoding="utf-8") as f:
                f.write("ERROR DE SUNAT: Los números de RUC ingresados no son válidos o no existen.\n")
            return save_path

        # Buscar otros posibles mensajes dentro de la clase error de SUNAT
        error_match = re.search(r'<p class="error">(.*?)</p>', response.text, flags=re.DOTALL)
        if error_match:
            err_msg = error_match.group(1).strip()
            # Reemplazar entidades HTML comunes
            err_msg = err_msg.replace("&amp;", "&").replace("&amp;", "&").replace("&oacute;", "ó").replace("&iacute;", "í").replace("&aacute;", "á").replace("&eacute;", "é").replace("&uacute;", "ú").replace("&ntilde;", "ñ")
            filename = f"resultado_sunat_batch_{batch_index}_ERROR.txt"
            save_path = os.path.normpath(os.path.join(output_dir, filename))
            with open(save_path, "w", encoding="utf-8") as f:
                f.write(f"ERROR DE SUNAT: {err_msg}\n")
            return save_path

    filename = get_filename_from_response(response) or f"resultado_sunat_batch_{batch_index}"
    _, ext = os.path.splitext(filename)
    ext = ext.lower()

    if not ext:
        ext = ".rar" if is_archive_content(response) else ".txt"
        filename += ext

    save_path = os.path.join(output_dir, filename)

    if is_archive_content(response):
        return save_binary_response(response, save_path)

    archive_from_html = download_archive_from_html(response, output_dir)
    if archive_from_html:
        return archive_from_html

    if ext == ".txt" or response.headers.get("content-type", "").lower().startswith("text"):
        save_path = os.path.normpath(save_path)
        with open(save_path, "w", encoding="utf-8", errors="replace") as f:
            f.write(response.text)
        return save_path

    return save_response_to_excel(response.text, os.path.join(output_dir, f"resultado_sunat_batch_{batch_index}.xlsx"))


def load_rucs_from_file(file_path: str) -> List[str]:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"No existe el archivo: {file_path}")

    if file_path.lower().endswith(".txt"):
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            values = [normalize_ruc_text(line) for line in f if line.strip()]
    elif file_path.lower().endswith((".xls", ".xlsx")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError(
                "Para abrir archivos Excel es necesario instalar openpyxl: python -m pip install openpyxl"
            ) from exc
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
        sheet = workbook.active
        values = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=1, values_only=True), start=1):
            candidate = extract_document_from_row(row, row_index)
            if candidate:
                values.append(candidate)
    else:
        raise ValueError("El archivo debe ser .txt, .xls o .xlsx")

    values = [v for v in values if v]
    if not values:
        raise ValueError("El archivo no contiene ningún RUC válido.")
    return values


def consult_ruc_manual(rucs: List[str]) -> requests.Response:
    if not rucs:
        raise ValueError("Debe indicar al menos un RUC para consultar.")

    session = requests.Session()
    session.headers.update(HEADERS)

    response = session.get(BASE_URL, timeout=30)
    response.raise_for_status()

    token = generate_sunat_token(52)
    payload = [("accion", "consManual"), ("token", token), ("txtRuc", rucs[0])]
    for ruc in rucs:
        payload.append(("selRuc", ruc))

    response = session.post(BASE_URL, data=payload, timeout=60)
    response.raise_for_status()
    return response


def chunk_values(values: List[str], batch_size: int = 100) -> List[List[str]]:
    return [values[i : i + batch_size] for i in range(0, len(values), batch_size)]


def get_output_dir(file_path: str) -> str:
    base_dir = os.path.dirname(file_path) or os.getcwd()
    output_dir = os.path.join(base_dir, "resultados_sunat")
    if os.path.isfile(output_dir):
        output_dir = os.path.join(base_dir, f"resultados_sunat_{random.randint(1000,9999)}")
    try:
        os.makedirs(output_dir, exist_ok=True)
        return output_dir
    except OSError:
        fallback_dir = os.path.join(os.getcwd(), "resultados_sunat")
        os.makedirs(fallback_dir, exist_ok=True)
        return fallback_dir

def clean_val(val: Optional[str], fallback: str = "N/A") -> str:
    if val is None:
        return fallback
    v = str(val).strip()
    if not v or v == "-" or v.upper() == "NONE" or v.upper() == "N/A":
        return fallback
    return v


def consolidate_results(output_dir: str) -> Optional[str]:
    excel_path = os.path.join(output_dir, "resultados_consolidados.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"

    # Escribir cabeceras solicitadas por el usuario
    headers = [
        "NumDocumento", "TipoDocumento", "Apellido_Paterno", "Apellido_Materno",
        "Nombres", "Direccion", "EstadoCivil", "Ubigeo", "Restriccion"
    ]
    ws.append(headers)

    has_data = False

    # Recorrer todos los archivos zip en la carpeta de salida
    for file_name in sorted(os.listdir(output_dir)):
        if not file_name.lower().endswith(".zip"):
            continue

        zip_path = os.path.join(output_dir, file_name)
        try:
            with zipfile.ZipFile(zip_path, "r") as zf:
                for name in zf.namelist():
                    if name.lower().endswith(".txt") and not name.startswith("__"):
                        content = zf.read(name).decode("utf-8", errors="replace")
                        rows = parse_pipe_delimited_text(content)
                        if not rows:
                            continue

                        start_row_idx = 0
                        if rows and "NumeroRuc" in rows[0][0]:
                            start_row_idx = 1

                        for row in rows[start_row_idx:]:
                            if len(row) < 2:
                                continue

                            num_doc = clean_val(row[0], fallback="") if len(row) > 0 else ""
                            if not num_doc:
                                continue

                            nombres = clean_val(row[1], fallback="N/A") if len(row) > 1 else "N/A"
                            estado = clean_val(row[6], fallback="N/A") if len(row) > 6 else "N/A"
                            dep = clean_val(row[9], fallback="") if len(row) > 9 else ""
                            prov = clean_val(row[10], fallback="") if len(row) > 10 else ""
                            dist = clean_val(row[11], fallback="") if len(row) > 11 else ""
                            direccion = clean_val(row[12], fallback="N/A") if len(row) > 12 else "N/A"

                            if dep or prov or dist:
                                ubigeo = f"{dep or 'N/A'}/{prov or 'N/A'}/{dist or 'N/A'}"
                            else:
                                ubigeo = "N/A"

                            new_row = [
                                num_doc,                     # NumDocumento
                                "RUC",                       # TipoDocumento
                                "N/A",                       # Apellido_Paterno
                                "N/A",                       # Apellido_Materno
                                nombres,                     # Nombres
                                direccion,                   # Direccion
                                "N/A",                       # EstadoCivil
                                ubigeo,                      # Ubigeo
                                estado                       # Restriccion
                            ]
                            ws.append(new_row)
                            has_data = True
        except Exception as e:
            print(f"Error consolidando archivo {file_name}: {e}")

    if has_data:
        # Auto-ajustar el ancho de las columnas de Excel
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        try:
            wb.save(excel_path)
            return excel_path
        except PermissionError:
            # Fallback en caso de que esté abierto el archivo
            fallback_path = os.path.join(output_dir, f"resultados_consolidados_{random.randint(1000,9999)}.xlsx")
            wb.save(fallback_path)
            return fallback_path
    return None


# ── Consolidado DNI + RUC ─────────────────────────────────────────────────────

_COL_ALIAS: dict = {
    "DNI":              ["DNI", "NUMDNI", "NUM_DNI", "NUMDOCUMENTO", "NUM_DOCUMENTO", "DOCUMENTO"],
    "APELLIDO_PATERNO": ["APELLIDO_PATERNO", "APELLIDOPATERNO", "APE_PATERNO", "APEPATERNO"],
    "APELLIDO_MATERNO": ["APELLIDO_MATERNO", "APELLIDOMATERNO", "APE_MATERNO", "APEMATERNO"],
    "NOMBRES":          ["NOMBRES", "NOMBRE", "NAMES", "NAME"],
    "ESTADOCIVIL":      ["ESTADOCIVIL", "ESTADOCIVIL", "ESTADO_CIVIL"],
    "SEXO":             ["SEXO", "GENERO", "GENDER"],
    "DIRECCION":        ["DIRECCION", "DIRECCION", "ADDRESS"],
    "UBIGEO":           ["UBIGEO", "COD_UBIGEO", "CODUBIGEO"],
    "RESTRICCION":      ["RESTRICCION", "RESTRICCION"],
}


def _norm_col(s: str) -> str:
    s = s.upper().strip()
    for src, dst in [("Á","A"),("É","E"),("Í","I"),("Ó","O"),("Ú","U"),("Ñ","N"),
                     ("á","A"),("é","E"),("í","I"),("ó","O"),("ú","U"),("ñ","N")]:
        s = s.replace(src, dst)
    return re.sub(r"[\s_\-]", "", s)


def _find_col(header: list, field: str) -> int:
    aliases = [_norm_col(a) for a in _COL_ALIAS.get(field, [field])]
    for i, h in enumerate(header):
        if _norm_col(str(h)) in aliases:
            return i
    return -1


def load_input_excel_data(file_path: str) -> List[dict]:
    """Read the DNI-based input Excel and return a list of row dicts."""
    try:
        from openpyxl import load_workbook as _lw
    except ImportError as exc:
        raise ImportError("Instale openpyxl: python -m pip install openpyxl") from exc

    wb = _lw(filename=file_path, read_only=True, data_only=True)
    ws = wb.active

    col_map: dict = {}
    rows_data: List[dict] = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [str(c).strip() if c is not None else "" for c in row]

        if row_idx == 1:
            for field in _COL_ALIAS:
                idx = _find_col(cells, field)
                if idx >= 0:
                    col_map[field] = idx
            continue

        if not any(cells):
            continue

        _NULL_VALS = {"null", "none", "n/a", "-", ""}

        def _get(field: str, _cells=cells) -> str:
            idx = col_map.get(field, -1)
            if 0 <= idx < len(_cells):
                v = _cells[idx].strip()
                return "N/A" if v.lower() in _NULL_VALS else v
            return "N/A"

        dni_idx = col_map.get("DNI", -1)
        dni_raw = cells[dni_idx] if 0 <= dni_idx < len(cells) else ""
        # Handle Excel float representation: "76833382.0" → "76833382"
        if "." in dni_raw:
            try:
                dni_clean = str(int(float(dni_raw)))
            except (ValueError, OverflowError):
                dni_clean = re.sub(r"[^0-9]", "", dni_raw.split(".")[0])
        else:
            dni_clean = re.sub(r"[^0-9]", "", dni_raw)
        if not dni_clean:
            continue

        rows_data.append({
            "DNI":              dni_clean,
            "APELLIDO_PATERNO": _get("APELLIDO_PATERNO"),
            "APELLIDO_MATERNO": _get("APELLIDO_MATERNO"),
            "NOMBRES":          _get("NOMBRES"),
            "ESTADOCIVIL":      _get("ESTADOCIVIL"),
            "SEXO":             _get("SEXO"),
            "DIRECCION":        _get("DIRECCION"),
            "UBIGEO":           _get("UBIGEO"),
            "RESTRICCION":      _get("RESTRICCION"),
        })

    if not rows_data:
        raise ValueError("No se encontraron filas con DNI válido en el archivo.")
    return rows_data


def parse_sunat_zip_results(output_dir: str) -> dict:
    """Parse SUNAT zip files and return dict: ruc_number -> {NOMBRES, ESTADO, DIRECCION, UBIGEO}."""
    result: dict = {}
    for fname in sorted(os.listdir(output_dir)):
        if not fname.lower().endswith(".zip"):
            continue
        try:
            with zipfile.ZipFile(os.path.join(output_dir, fname), "r") as zf:
                for inner in zf.namelist():
                    if not inner.lower().endswith(".txt") or inner.startswith("__"):
                        continue
                    content = zf.read(inner).decode("utf-8", errors="replace")
                    rows = parse_pipe_delimited_text(content)
                    if not rows:
                        continue
                    start = 1 if rows and "NumeroRuc" in rows[0][0] else 0
                    for row in rows[start:]:
                        if len(row) < 2:
                            continue
                        num_doc = row[0].strip().strip("|")
                        if len(num_doc) != 11 or not num_doc.isdigit():
                            continue
                        nombres   = clean_val(row[1])               if len(row) > 1  else "N/A"
                        estado    = clean_val(row[6])               if len(row) > 6  else "N/A"
                        dep       = clean_val(row[9],  fallback="") if len(row) > 9  else ""
                        prov      = clean_val(row[10], fallback="") if len(row) > 10 else ""
                        dist      = clean_val(row[11], fallback="") if len(row) > 11 else ""
                        direccion = clean_val(row[12])              if len(row) > 12 else "N/A"
                        ubigeo = f"{dep or 'N/A'}/{prov or 'N/A'}/{dist or 'N/A'}" if (dep or prov or dist) else "N/A"
                        result[num_doc] = {"NOMBRES": nombres, "ESTADO": estado, "DIRECCION": direccion, "UBIGEO": ubigeo}
        except Exception as exc:
            print(f"[parse_sunat] error en {fname}: {exc}")
    return result


def _auto_fit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 12)


def _save_wb(wb, path: str) -> str:
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = os.path.splitext(path)[0] + f"_{random.randint(1000,9999)}.xlsx"
        wb.save(alt)
        return alt


def build_consolidated_output(input_data: List[dict], sunat_data: dict, output_dir: str) -> str:
    """Excel en formato NumDocumento/TipoDocumento: primero DNIs (8 dig), luego RUCs (11 dig) de SUNAT."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    ws.append(["NumDocumento", "TipoDocumento", "Apellido_Paterno", "Apellido_Materno",
               "Nombres", "Direccion", "EstadoCivil", "Ubigeo", "Restriccion"])

    for r in (x for x in input_data if len(x["DNI"]) != 11):
        ws.append([r["DNI"], "DNI", r["APELLIDO_PATERNO"], r["APELLIDO_MATERNO"],
                   r["NOMBRES"], r["DIRECCION"], r["ESTADOCIVIL"], r["UBIGEO"], r["RESTRICCION"]])

    for r in (x for x in input_data if len(x["DNI"]) == 11):
        s = sunat_data.get(r["DNI"], {})
        ws.append([r["DNI"], "RUC", "N/A", "N/A",
                   s.get("NOMBRES", "N/A"), s.get("DIRECCION", "N/A"),
                   "N/A", s.get("UBIGEO", "N/A"), s.get("ESTADO", "N/A")])

    _auto_fit(ws)
    return _save_wb(wb, os.path.join(output_dir, "consolidado_final.xlsx"))


def build_updated_input_excel(input_data: List[dict], sunat_data: dict, source_path: str) -> str:
    """Excel en formato original del input: DNI rows sin tocar, RUC rows con datos de SUNAT."""
    base_dir = os.path.dirname(source_path)
    base_name = os.path.splitext(os.path.basename(source_path))[0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Actualizado"
    ws.append(["DNI", "Apellido_Paterno", "Apellido_Materno", "Nombres",
               "Direccion", "EstadoCivil", "Ubigeo", "Restriccion"])

    for r in (x for x in input_data if len(x["DNI"]) != 11):
        ws.append([r["DNI"], r["APELLIDO_PATERNO"], r["APELLIDO_MATERNO"], r["NOMBRES"],
                   r["DIRECCION"], r["ESTADOCIVIL"], r["UBIGEO"], r["RESTRICCION"]])

    for r in (x for x in input_data if len(x["DNI"]) == 11):
        s = sunat_data.get(r["DNI"], {})
        ws.append([r["DNI"], "N/A", "N/A",
                   s.get("NOMBRES", "N/A"),
                   s.get("DIRECCION", "N/A"), "N/A",
                   s.get("UBIGEO", "N/A"), s.get("ESTADO", "N/A")])

    _auto_fit(ws)
    return _save_wb(wb, os.path.join(base_dir, f"{base_name}_actualizado.xlsx"))


def process_full_consolidado(file_path: str, progress_callback=None) -> tuple:
    """Lee input Excel → consulta SUNAT solo para RUCs (11 dig) → genera 2 archivos de salida."""
    input_data = load_input_excel_data(file_path)
    ruc_to_query = [normalize_ruc_text(r["DNI"]) for r in input_data if len(r["DNI"]) == 11 and r["DNI"].isdigit()]
    total = len(ruc_to_query)
    output_dir = get_output_dir(file_path)

    if total == 0:
        if progress_callback:
            progress_callback("start", {"total_rucs": 0, "total_batches": 0})
            progress_callback("consolidate_start", {"output_dir": output_dir})
        sunat_data: dict = {}
        c = build_consolidated_output(input_data, sunat_data, output_dir)
        a = build_updated_input_excel(input_data, sunat_data, file_path)
        if progress_callback:
            progress_callback("consolidate_success", {"output_dir": output_dir, "consolidated_path": c, "actualizado_path": a})
        return c, a

    batches = chunk_values(ruc_to_query, batch_size=100)
    total_batches = len(batches)

    if progress_callback:
        progress_callback("start", {"total_rucs": total, "total_batches": total_batches})

    session = requests.Session()
    session.headers.update(HEADERS)

    for batch_index, batch in enumerate(batches, start=1):
        start_idx = (batch_index - 1) * 100 + 1
        end_idx = min(batch_index * 100, total)

        if progress_callback:
            progress_callback("batch_start", {
                "batch_index": batch_index, "total_batches": total_batches,
                "start_idx": start_idx, "end_idx": end_idx, "batch_len": len(batch),
            })

        with tempfile.NamedTemporaryFile("w", delete=False, suffix=".txt", encoding="utf-8") as tf:
            tf.writelines(f"{v}\n" for v in batch)
            temp_path = tf.name

        try:
            session.get(BASE_URL, timeout=30).raise_for_status()
            token = generate_sunat_token(52)
            payload = {"accion": "consArchivo", "tokenArch": token}
            with open(temp_path, "rb") as archivo:
                files = {"archivo": (os.path.basename(temp_path), archivo, "text/plain")}
                response = session.post(BASE_URL, data=payload, files=files, timeout=120)
            response.raise_for_status()
            saved = save_sunat_response(response, output_dir, batch_index)

            if progress_callback:
                if "_ERROR" in saved:
                    progress_callback("batch_warning", {
                        "batch_index": batch_index, "total_batches": total_batches,
                        "saved_path": saved, "warning_msg": "RUCs no válidos en este lote.",
                    })
                else:
                    progress_callback("batch_success", {
                        "batch_index": batch_index, "total_batches": total_batches,
                        "saved_path": saved, "status_code": response.status_code,
                    })
        except Exception as exc:
            if progress_callback:
                progress_callback("batch_error", {"batch_index": batch_index, "error_msg": str(exc)})
            raise
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    if progress_callback:
        progress_callback("consolidate_start", {"output_dir": output_dir})

    sunat_data = parse_sunat_zip_results(output_dir)
    consolidado = build_consolidated_output(input_data, sunat_data, output_dir)
    actualizado = build_updated_input_excel(input_data, sunat_data, file_path)

    if progress_callback:
        progress_callback("consolidate_success", {
            "output_dir": output_dir,
            "consolidated_path": consolidado,
            "actualizado_path": actualizado,
        })

    return consolidado, actualizado


def consult_ruc_file(file_path: str, progress_callback=None) -> List[tuple[int, str, requests.Response]]:
    values = load_rucs_from_file(file_path)
    total_rucs = len(values)
    batches = chunk_values(values, batch_size=100)
    total_batches = len(batches)
    output_dir = get_output_dir(file_path)
    results: List[tuple[int, str, requests.Response]] = []

    if progress_callback:
        progress_callback("start", {"total_rucs": total_rucs, "total_batches": total_batches})

    session = requests.Session()
    session.headers.update(HEADERS)

    for batch_index, batch in enumerate(batches, start=1):
        start_idx = (batch_index - 1) * 100 + 1
        end_idx = min(batch_index * 100, total_rucs)

        if progress_callback:
            progress_callback("batch_start", {
                "batch_index": batch_index,
                "total_batches": total_batches,
                "start_idx": start_idx,
                "end_idx": end_idx,
                "batch_len": len(batch)
            })

        with tempfile.NamedTemporaryFile("w", delete=False, suffix=".txt", encoding="utf-8") as temp_file:
            temp_file.writelines(f"{v}\n" for v in batch)
            temp_path = temp_file.name

        try:
            response = session.get(BASE_URL, timeout=30)
            response.raise_for_status()

            token = generate_sunat_token(52)
            payload = {"accion": "consArchivo", "tokenArch": token}
            with open(temp_path, "rb") as archivo:
                files = {"archivo": (os.path.basename(temp_path), archivo, "text/plain")}
                response = session.post(BASE_URL, data=payload, files=files, timeout=120)

            response.raise_for_status()
            saved_path = save_sunat_response(response, output_dir, batch_index)
            results.append((batch_index, saved_path, response))

            if progress_callback:
                if "_ERROR" in saved_path:
                    # Se detectó un error en el lote, pero continuamos con los demás lotes
                    progress_callback("batch_warning", {
                        "batch_index": batch_index,
                        "total_batches": total_batches,
                        "saved_path": saved_path,
                        "warning_msg": "Los RUCs de este lote no son válidos o no existen."
                    })
                else:
                    progress_callback("batch_success", {
                        "batch_index": batch_index,
                        "total_batches": total_batches,
                        "saved_path": saved_path,
                        "status_code": response.status_code
                    })
        except Exception as exc:
            if progress_callback:
                progress_callback("batch_error", {
                    "batch_index": batch_index,
                    "error_msg": str(exc)
                })
            raise
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    # Consolidar resultados
    if progress_callback:
        progress_callback("consolidate_start", {"output_dir": output_dir})

    consolidate_results(output_dir)

    if progress_callback:
        progress_callback("consolidate_success", {
            "output_dir": output_dir,
            "consolidated_path": os.path.normpath(os.path.join(output_dir, "resultados_consolidados.xlsx"))
        })

    return results


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Consulta un RUC en SUNAT desde la web de consulta múltiple."
    )
    parser.add_argument("ruc", nargs="*", help="Uno o más números de RUC para consulta manual.")
    parser.add_argument("--file", "-f", help="Archivo .txt o Excel (.xls/.xlsx) con RUCs.")
    parser.add_argument("--save", help="Guardar el HTML de respuesta en un archivo.")
    parser.add_argument("--gui", action="store_true", help="Abrir la interfaz Tkinter para seleccionar el archivo.")
    return parser.parse_args()


def print_response_summary(response: requests.Response, save_path: Optional[str] = None) -> None:
    print(f"HTTP {response.status_code}")
    if save_path:
        with open(save_path, "w", encoding="utf-8") as f:
            f.write(response.text)
        print(f"Respuesta guardada en {save_path}")
        return

    text = response.text
    print(text[:4000])
    if len(text) > 4000:
        print("\n... salida truncada, use --save para guardar el resultado completo.")


def print_batch_summary(results: List[tuple[int, str, requests.Response]]) -> None:
    for batch_index, save_path, response in results:
        print(f"Batch {batch_index}: HTTP {response.status_code} -> {save_path}")


def get_default_file_dir() -> str:
    home_dir = os.path.expanduser("~")
    dni_dir = os.path.join(home_dir, "Desktop", "DNI")
    if os.path.isdir(dni_dir):
        return dni_dir
    desktop_dir = os.path.join(home_dir, "Desktop")
    if os.path.isdir(desktop_dir):
        return desktop_dir
    downloads_dir = os.path.join(home_dir, "Downloads")
    if os.path.isdir(downloads_dir):
        return downloads_dir
    return home_dir


def open_file_dialog() -> Optional[str]:
    initial_dir = get_default_file_dir()
    return filedialog.askopenfilename(
        initialdir=initial_dir,
        title="Seleccionar archivo de RUC",
        filetypes=[("Text files", "*.txt"), ("Excel files", "*.xlsx *.xls"), ("Todos los archivos", "*")],
    )


def run_gui() -> None:
    if tk is None:
        print("No se puede ejecutar la GUI porque Tkinter no está disponible.")
        return

    # Design system — 8px grid
    C_BG      = "#0F172A"
    C_SURFACE = "#1E293B"
    C_BORDER  = "#334155"
    C_ACCENT  = "#0EA5E9"
    C_ACCENT_H= "#0284C7"
    C_SUCCESS = "#10B981"
    C_WARNING = "#F59E0B"
    C_ERROR   = "#EF4444"
    C_TEXT    = "#F1F5F9"
    C_MUTED   = "#64748B"
    C_INPUT   = "#0B1120"
    F         = "Segoe UI"

    root = tk.Tk()
    root.title("SUNAT RUC Pro")
    root.geometry("800x580")
    root.configure(bg=C_BG)
    root.minsize(640, 480)

    selected_path = tk.StringVar()
    status_var    = tk.StringVar(value="Listo para iniciar")

    def _log(msg: str, tag: str = "") -> None:
        log_text.config(state=tk.NORMAL)
        if tag:
            log_text.insert(tk.END, msg, tag)
        else:
            log_text.insert(tk.END, msg)
        log_text.see(tk.END)
        log_text.config(state=tk.DISABLED)

    def _status(msg: str, color: str = C_MUTED) -> None:
        status_var.set(msg)
        status_dot.config(fg=color)
        status_lbl.config(fg=color)

    def select_file() -> None:
        path = open_file_dialog()
        if path:
            selected_path.set(path)
            _log(f"Archivo: {path}\n")

    def send_file() -> None:
        path = selected_path.get()
        if not path:
            messagebox.showwarning("Sin archivo", "Seleccione un archivo .txt o Excel primero.")
            return

        btn_select.config(state=tk.DISABLED, bg=C_BORDER, fg=C_MUTED)
        btn_send.config(state=tk.DISABLED, bg=C_BORDER, fg=C_MUTED)

        def _worker():
            def cb(evt, data):
                if evt == "start":
                    _log(f"\n{data['total_rucs']} RUCs encontrados — {data['total_batches']} lotes\n")
                    _status(f"Procesando 0 / {data['total_batches']} lotes…", C_ACCENT)
                elif evt == "batch_start":
                    _log(f"→  Lote {data['batch_index']}/{data['total_batches']}  (RUC {data['start_idx']} – {data['end_idx']})\n")
                    _status(f"Lote {data['batch_index']} / {data['total_batches']}", C_ACCENT)
                elif evt == "batch_success":
                    _log(f"   ✓  {data['saved_path']}\n", "success")
                    _status(f"Lote {data['batch_index']} completado", C_SUCCESS)
                elif evt == "batch_warning":
                    _log(f"   ⚠  RUCs inválidos  →  {data['saved_path']}\n", "warning")
                    _status(f"Lote {data['batch_index']} con observación", C_WARNING)
                elif evt == "batch_error":
                    _log(f"   ✗  Lote {data['batch_index']}: {data['error_msg']}\n", "error")
                    _status("Error en la consulta", C_ERROR)
                elif evt == "consolidate_start":
                    _log("\nConsolidando resultados…\n")
                    _status("Consolidando en Excel…", C_ACCENT)
                elif evt == "consolidate_success":
                    _log(f"   ✓  {data['consolidated_path']}\n", "success")
                    _status("Consolidado listo", C_SUCCESS)

            try:
                _log("\nIniciando consulta masiva en SUNAT…\n")
                results = consult_ruc_file(path, progress_callback=cb)
                _status("Consulta completada", C_SUCCESS)
                messagebox.showinfo(
                    "Completado",
                    f"Lotes procesados: {len(results)}\n\n"
                    f"Resultados en:\n"
                    f"{os.path.join(os.path.dirname(path), 'resultados_sunat', 'resultados_consolidados.xlsx')}"
                )
            except Exception as exc:
                _log(f"\nError: {exc}\n", "error")
                _status("Error en la consulta", C_ERROR)
                messagebox.showerror("Error", str(exc))
            finally:
                btn_select.config(state=tk.NORMAL, bg=C_SURFACE, fg=C_TEXT)
                btn_send.config(state=tk.NORMAL, bg=C_ACCENT, fg="#0F172A")

        threading.Thread(target=_worker, daemon=True).start()

    def generar_consolidado() -> None:
        path = open_file_dialog()
        if not path:
            return
        selected_path.set(path)
        _log(f"Archivo: {path}\n")

        btn_select.config(state=tk.DISABLED, bg=C_BORDER, fg=C_MUTED)
        btn_send.config(state=tk.DISABLED, bg=C_BORDER, fg=C_MUTED)
        btn_consolidar.config(state=tk.DISABLED, bg=C_BORDER, fg=C_MUTED)

        def _worker_cons():
            def cb(evt, data):
                if evt == "start":
                    _log(f"\n{data['total_rucs']} documentos encontrados — {data['total_batches']} lotes\n")
                    _status(f"Consultando SUNAT…", C_ACCENT)
                elif evt == "batch_start":
                    _log(f"→  Lote {data['batch_index']}/{data['total_batches']}  (doc {data['start_idx']} – {data['end_idx']})\n")
                    _status(f"Lote {data['batch_index']} / {data['total_batches']}", C_ACCENT)
                elif evt == "batch_success":
                    _log(f"   ✓  {data['saved_path']}\n", "success")
                    _status(f"Lote {data['batch_index']} completado", C_SUCCESS)
                elif evt == "batch_warning":
                    _log(f"   ⚠  {data['saved_path']}\n", "warning")
                    _status(f"Lote {data['batch_index']} con observación", C_WARNING)
                elif evt == "batch_error":
                    _log(f"   ✗  Error lote {data['batch_index']}: {data['error_msg']}\n", "error")
                    _status("Error en consulta", C_ERROR)
                elif evt == "consolidate_start":
                    _log("\nGenerando consolidado final…\n")
                    _status("Generando Excel consolidado…", C_ACCENT)
                elif evt == "consolidate_success":
                    _log(f"   ✓  Consolidado: {data['consolidated_path']}\n", "success")
                    _log(f"   ✓  Actualizado: {data['actualizado_path']}\n", "success")
                    _status("Consolidado listo", C_SUCCESS)

            try:
                _log("\nIniciando consolidado DNI + RUC…\n")
                consolidado, actualizado = process_full_consolidado(path, progress_callback=cb)
                _status("Consolidado generado", C_SUCCESS)
                messagebox.showinfo(
                    "Consolidado listo",
                    f"Archivos generados:\n\n"
                    f"1. Consolidado (nuevo formato):\n{consolidado}\n\n"
                    f"2. Actualizado (formato original):\n{actualizado}"
                )
            except Exception as exc:
                _log(f"\nError: {exc}\n", "error")
                _status("Error en el proceso", C_ERROR)
                messagebox.showerror("Error", str(exc))
            finally:
                btn_select.config(state=tk.NORMAL, bg=C_SURFACE, fg=C_TEXT)
                btn_send.config(state=tk.NORMAL, bg=C_ACCENT, fg="#0F172A")
                btn_consolidar.config(state=tk.NORMAL, bg=C_SUCCESS, fg="#0F172A")

        threading.Thread(target=_worker_cons, daemon=True).start()

    # ── Layout ────────────────────────────────────────────────────────
    pad = tk.Frame(root, bg=C_BG)
    pad.pack(fill=tk.BOTH, expand=True, padx=24, pady=24)

    # Header
    tk.Label(pad, text="SUNAT RUC Pro", font=(F, 20, "bold"),
             bg=C_BG, fg=C_TEXT).pack(anchor=tk.W)
    tk.Label(pad, text="Consulta masiva de RUC  ·  SUNAT Perú",
             font=(F, 10), bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(2, 16))

    tk.Frame(pad, height=1, bg=C_BORDER).pack(fill=tk.X, pady=(0, 16))

    # File section
    tk.Label(pad, text="ARCHIVO DE ENTRADA", font=(F, 8, "bold"),
             bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(0, 8))

    input_wrap = tk.Frame(pad, bg=C_INPUT, padx=12, pady=8)
    input_wrap.pack(fill=tk.X, pady=(0, 8))

    tk.Entry(
        input_wrap, textvariable=selected_path,
        font=(F, 10), bg=C_INPUT, fg=C_TEXT,
        insertbackground=C_TEXT, relief=tk.FLAT, bd=0
    ).pack(fill=tk.X)

    # Buttons
    btns = tk.Frame(pad, bg=C_BG)
    btns.pack(fill=tk.X, pady=(0, 16))

    btn_select = tk.Button(
        btns, text="Seleccionar archivo",
        font=(F, 10), bg=C_SURFACE, fg=C_TEXT,
        activebackground=C_BORDER, activeforeground=C_TEXT,
        relief=tk.FLAT, bd=0, padx=16, pady=8,
        cursor="hand2", command=select_file
    )
    btn_select.pack(side=tk.LEFT)

    btn_send = tk.Button(
        btns, text="Enviar a SUNAT",
        font=(F, 10, "bold"), bg=C_ACCENT, fg="#0F172A",
        activebackground=C_ACCENT_H, activeforeground="#0F172A",
        relief=tk.FLAT, bd=0, padx=16, pady=8,
        cursor="hand2", command=send_file
    )
    btn_send.pack(side=tk.LEFT, padx=(8, 0))

    btn_consolidar = tk.Button(
        btns, text="Generar Consolidado",
        font=(F, 10, "bold"), bg=C_SUCCESS, fg="#0F172A",
        activebackground="#059669", activeforeground="#0F172A",
        relief=tk.FLAT, bd=0, padx=16, pady=8,
        cursor="hand2", command=generar_consolidado
    )
    btn_consolidar.pack(side=tk.LEFT, padx=(8, 0))

    def _hover(btn, on_bg, off_bg):
        btn.bind("<Enter>", lambda _: btn.config(bg=on_bg) if btn["state"] == tk.NORMAL else None)
        btn.bind("<Leave>", lambda _: btn.config(bg=off_bg) if btn["state"] == tk.NORMAL else None)

    _hover(btn_select, C_BORDER, C_SURFACE)
    _hover(btn_send, C_ACCENT_H, C_ACCENT)
    _hover(btn_consolidar, "#059669", C_SUCCESS)

    # Status
    tk.Frame(pad, height=1, bg=C_BORDER).pack(fill=tk.X, pady=(0, 8))

    status_row = tk.Frame(pad, bg=C_BG)
    status_row.pack(fill=tk.X, pady=(0, 8))

    status_dot = tk.Label(status_row, text="●", font=(F, 9), bg=C_BG, fg=C_MUTED)
    status_dot.pack(side=tk.LEFT)
    status_lbl = tk.Label(status_row, textvariable=status_var,
                          font=(F, 9), bg=C_BG, fg=C_MUTED, anchor=tk.W)
    status_lbl.pack(side=tk.LEFT, padx=(4, 0))

    # Log
    tk.Label(pad, text="BITÁCORA", font=(F, 8, "bold"),
             bg=C_BG, fg=C_MUTED).pack(anchor=tk.W, pady=(0, 4))

    log_text = scrolledtext.ScrolledText(
        pad, font=("Consolas", 9),
        bg=C_INPUT, fg="#94A3B8",
        relief=tk.FLAT, bd=0,
        insertbackground=C_TEXT,
        padx=12, pady=8
    )
    log_text.pack(fill=tk.BOTH, expand=True)

    log_text.tag_config("success", foreground=C_SUCCESS)
    log_text.tag_config("warning", foreground=C_WARNING)
    log_text.tag_config("error",   foreground=C_ERROR)

    _log("Consola activa. Seleccione un archivo y presione 'Enviar a SUNAT'.\n")

    root.mainloop()


def main() -> int:
    args = parse_args()
    if args.gui or (not args.file and not args.ruc):
        run_gui()
        return 0

    if args.file and args.ruc:
        print("Error: use --file para carga de archivo o pase RUCs manuales, no ambas opciones.", file=sys.stderr)
        return 1

    try:
        if args.file:
            def cli_callback(event_type, data):
                if event_type == "start":
                    print(f"\n[INFO] Se encontraron {data['total_rucs']} RUCs válidos en el archivo.")
                    print(f"[INFO] Procesando en {data['total_batches']} lotes de 100...")
                elif event_type == "batch_start":
                    print(f"-> Lote {data['batch_index']}/{data['total_batches']}: Consultando RUCs del {data['start_idx']} al {data['end_idx']}...", end="", flush=True)
                elif event_type == "batch_success":
                    print(f" completado -> {data['saved_path']}")
                elif event_type == "batch_warning":
                    print(f" OBSERVACIÓN -> {data['saved_path']} ({data['warning_msg']})")
                elif event_type == "batch_error":
                    print(f" ERROR: {data['error_msg']}")
                elif event_type == "consolidate_start":
                    print(f"\n[INFO] Consolidando resultados en Excel único...")
                elif event_type == "consolidate_success":
                    print(f"[OK] Excel consolidado creado exitosamente en: {data['consolidated_path']}")

            results = consult_ruc_file(args.file, progress_callback=cli_callback)
            print(f"\n[INFO] Consulta completada. Se generaron {len(results)} archivo(s) de resultado y se unificaron en resultados_consolidados.xlsx.")
            if args.save:
                print("Advertencia: --save no aplica al enviar un archivo. Se han guardado los resultados en la carpeta resultados_sunat.")
            return 0
        else:
            response = consult_ruc_manual(args.ruc)
    except requests.RequestException as exc:
        print(f"Error de conexión: {exc}", file=sys.stderr)
        return 1
    except (ValueError, FileNotFoundError, ImportError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    print_response_summary(response, args.save)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
